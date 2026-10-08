# Copyright 2026 Omnifloo
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl.html).

import io
import json

from openpyxl import load_workbook

from odoo.tests import HttpCase, tagged

XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@tagged("post_install", "-at_install")
class TestReportXlsxController(HttpCase):
    """Exercise the /report/download route overridden by report_xlsx.

    The download goes through ``content_disposition`` and the error path through
    ``serialize_exception``: both helpers moved in Odoo 20 (``odoo.http.stream``
    and ``odoo.http.dispatcher``), which only an HTTP request can catch.
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.report = cls.env["ir.actions.report"].create(
            {
                "name": "Print to XLSX",
                "model": "res.partner",
                "report_type": "xlsx",
                "report_name": "report_xlsx.partner_xlsx",
            }
        )
        cls.partner = cls.env["res.partner"].create({"name": "XLSX Controller Partner"})

    def _download(self, url):
        self.authenticate("admin", "admin")
        return self.url_open(
            "/report/download",
            data={
                "data": json.dumps([url, "xlsx"]),
                "context": json.dumps({}),
                "csrf_token": self.csrf_token(),
            },
        )

    def test_download_xlsx(self):
        response = self._download(
            f"/report/xlsx/report_xlsx.partner_xlsx/{self.partner.id}"
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["Content-Type"], XLSX_MIMETYPE)
        disposition = response.headers.get("Content-Disposition", "")
        self.assertIn("attachment", disposition)
        self.assertIn("Print%20to%20XLSX.xlsx", disposition)
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.cell(1, 1).value, self.partner.name)

    def test_download_xlsx_error(self):
        with self.assertLogs("odoo.addons.report_xlsx.controllers.main", "ERROR"):
            response = self._download("/report/xlsx/report_xlsx.no_such_report/1")
        self.assertEqual(response.status_code, 200)
        self.assertNotIn(XLSX_MIMETYPE, response.headers.get("Content-Type", ""))
        error = json.loads(response.text.replace("&#34;", '"').replace("&quot;", '"'))
        self.assertEqual(error["message"], "Odoo Server Error")
        self.assertIn("name", error["data"])
